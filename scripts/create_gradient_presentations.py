"""Create single-Z timecourse and optional Step 8 all-Z presentations."""

from __future__ import annotations

import argparse
import io
import json
import re
from pathlib import Path
import sys

from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
GREEN = RGBColor(32, 168, 62)
MAGENTA = RGBColor(214, 42, 139)
DARK = RGBColor(35, 38, 43)
GRAY = RGBColor(95, 100, 108)


def _new_presentation() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    return presentation


def _blank_slide(presentation: Presentation):
    return presentation.slides.add_slide(presentation.slide_layouts[6])


def _text(slide, text: str, x: float, y: float, width: float, height: float,
          *, size: int = 24, color: RGBColor = DARK, bold: bool = False,
          align: PP_ALIGN = PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def _title_slide(presentation: Presentation, title: str, subtitle: str) -> None:
    slide = _blank_slide(presentation)
    _text(slide, title, 0.8, 2.25, 11.7, 1.0, size=30, bold=True, align=PP_ALIGN.CENTER)
    _text(slide, subtitle, 1.0, 3.35, 11.3, 0.8, size=18, color=GRAY, align=PP_ALIGN.CENTER)


def _section_slide(presentation: Presentation, title: str, subtitle: str = "") -> None:
    slide = _blank_slide(presentation)
    _text(slide, title, 0.8, 2.45, 11.7, 0.8, size=30, bold=True, align=PP_ALIGN.CENTER)
    if subtitle:
        _text(slide, subtitle, 1.0, 3.4, 11.3, 0.7, size=17, color=GRAY, align=PP_ALIGN.CENTER)


def _image_dimensions(path: Path) -> tuple[int, int]:
    with Image.open(path) as image:
        return image.size


def _fit_box(path: Path, x: float, y: float, width: float, height: float) -> tuple[float, float, float, float]:
    image_width, image_height = _image_dimensions(path)
    scale = min(width / image_width, height / image_height)
    placed_width = image_width * scale
    placed_height = image_height * scale
    return x + (width - placed_width) / 2, y + (height - placed_height) / 2, placed_width, placed_height


def _compressed_picture(path: Path, max_width: int = 2200) -> io.BytesIO:
    with Image.open(path) as image:
        image = image.convert("RGB")
        if image.width > max_width:
            height = max(1, round(image.height * max_width / image.width))
            image = image.resize((max_width, height), Image.Resampling.LANCZOS)
        stream = io.BytesIO()
        image.save(stream, format="JPEG", quality=90, optimize=True)
    stream.seek(0)
    return stream


def _add_picture(slide, path: Path, x: float, y: float, width: float, height: float,
                 *, compress: bool = False) -> None:
    placed_x, placed_y, placed_width, placed_height = _fit_box(path, x, y, width, height)
    source = _compressed_picture(path) if compress else str(path)
    slide.shapes.add_picture(
        source,
        Inches(placed_x),
        Inches(placed_y),
        width=Inches(placed_width),
        height=Inches(placed_height),
    )


def _time_label(timepoint: int) -> str:
    return f"t = {timepoint * 2}  (t{timepoint:03d})"


def _indexed_files(folder: Path, pattern: str) -> dict[int, Path]:
    result: dict[int, Path] = {}
    expression = re.compile(pattern)
    for path in folder.glob("*.png"):
        match = expression.search(path.name)
        if match:
            result[int(match.group(1))] = path
    return result


def _run_z_index(run_dir: Path, indexed_paths: dict[int, Path] | None = None) -> int:
    metadata_path = run_dir / "run_metadata.json"
    if metadata_path.is_file():
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        return int(metadata["z_index_requested"])
    for path in (indexed_paths or {}).values():
        match = re.search(r"_z(\d{2})_", path.name)
        if match:
            return int(match.group(1))
    raise ValueError("Cannot determine the analyzed Z plane.")


def _indexed_step1_files(run_dir: Path, suffix: str) -> dict[int, Path]:
    result: dict[int, Path] = {}
    for path in (run_dir / "step_01_stitched_images").glob(f"t[0-9][0-9][0-9]/*{suffix}"):
        match = re.search(r"_t(\d{3})_", path.name)
        if match:
            result[int(match.group(1))] = path
    return result


def _single_z_image_group_slides(
    presentation: Presentation,
    label: str,
    paths: dict[int, Path],
    color: RGBColor,
    z_index: int,
) -> None:
    timepoints = sorted(paths)
    for offset in range(0, len(timepoints), 5):
        chunk = timepoints[offset:offset + 5]
        slide = _blank_slide(presentation)
        start, end = chunk[0] * 2, chunk[-1] * 2
        _text(
            slide,
            f"Z={z_index} · {label} · t={start}–{end}",
            0.35,
            0.12,
            12.65,
            0.45,
            size=21,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        row_height = 1.28
        for index, timepoint in enumerate(chunk):
            y = 0.72 + index * row_height
            _text(
                slide,
                f"t={timepoint * 2}",
                0.12,
                y + 0.32,
                0.75,
                0.35,
                size=13,
                color=GRAY,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
            _add_picture(slide, paths[timepoint], 0.92, y, 12.1, 1.05, compress=True)


def create_single_z_image_presentation(run_dir: Path) -> Path:
    groups = (
        ("GFP–Cy5 merge", _indexed_step1_files(run_dir, "_GFP-Cy5_merge.png"), DARK),
        ("GFP", _indexed_step1_files(run_dir, "_GFP_positions_left-to-right_preview.png"), GREEN),
        ("Cy5", _indexed_step1_files(run_dir, "_Cy5_positions_left-to-right_preview.png"), MAGENTA),
    )
    if not groups[0][1]:
        raise ValueError("No Step 1 timepoint images were found.")
    expected = set(groups[0][1])
    for label, paths, _ in groups:
        if set(paths) != expected:
            raise ValueError(f"Step 1 {label} images do not cover the same timepoints.")
    z_index = _run_z_index(run_dir, groups[0][1])
    presentation = _new_presentation()
    _title_slide(
        presentation,
        "Gradient images across the timecourse",
        f"Z{z_index} · five timepoints per slide · experimental time = ND2 timepoint × 2",
    )
    for label, paths, color in groups:
        _section_slide(presentation, label, f"Z{z_index} · {len(paths)} timepoints")
        _single_z_image_group_slides(presentation, label, paths, color, z_index)
    output = run_dir / f"2026_07_30_z{z_index:02d}_gradient_images_timecourse_5-per-slide_t2.pptx"
    presentation.save(output)
    return output


def create_single_z_comparison(run_dir: Path) -> Path:
    before = _indexed_files(
        run_dir / "step_02_feathered_profiles_before_correction",
        r"_t(\d{3})_.*before-correction\.png$",
    )
    after = _indexed_files(
        run_dir / "step_06_tanh_gradient_fits",
        r"_t(\d{3})_.*tanh-fit\.png$",
    )
    timepoints = sorted(set(before) & set(after))
    if not timepoints:
        raise ValueError("No matching Step 2 and Step 6 timepoints were found.")
    z_index = _run_z_index(run_dir, before)
    presentation = _new_presentation()
    _title_slide(
        presentation,
        "Profiles before correction and after correction with tanh fit",
        f"Z{z_index} · one slide per timepoint · experimental time = ND2 timepoint × 2",
    )
    for timepoint in timepoints:
        slide = _blank_slide(presentation)
        _text(slide, _time_label(timepoint), 0.4, 0.15, 12.5, 0.45, size=23, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, "Before correction · Step 02", 0.35, 0.67, 6.25, 0.35, size=16, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, "After correction + tanh fit · Step 06", 6.73, 0.67, 6.25, 0.35, size=16, color=MAGENTA, bold=True, align=PP_ALIGN.CENTER)
        _add_picture(slide, before[timepoint], 0.25, 1.05, 6.35, 6.15)
        _add_picture(slide, after[timepoint], 6.73, 1.05, 6.35, 6.15)
    output = run_dir / f"2026_07_30_z{z_index:02d}_profiles_before_vs_after_tanh_fit.pptx"
    presentation.save(output)
    return output


def create_z15_comparison(run_dir: Path) -> Path:
    """Backward-compatible alias for callers of the original helper."""
    return create_single_z_comparison(run_dir)


def _z_number(path: Path) -> int:
    match = re.search(r"_z(\d{2})_", path.name)
    if not match:
        raise ValueError(f"Cannot read Z number from {path.name}")
    return int(match.group(1))


def _stitched_group_slides(presentation: Presentation, timepoint: int, label: str,
                           paths: list[Path], color: RGBColor) -> None:
    paths = sorted(paths, key=_z_number)
    for offset in range(0, len(paths), 5):
        chunk = paths[offset:offset + 5]
        slide = _blank_slide(presentation)
        start_z, end_z = _z_number(chunk[0]), _z_number(chunk[-1])
        _text(slide, f"{_time_label(timepoint)} · {label} · Z{start_z:02d}–Z{end_z:02d}",
              0.35, 0.12, 12.65, 0.45, size=21, color=color, bold=True, align=PP_ALIGN.CENTER)
        row_height = 1.28
        for index, path in enumerate(chunk):
            y = 0.72 + index * row_height
            _text(slide, f"Z={_z_number(path)}", 0.18, y + 0.33, 0.65, 0.35, size=13, color=GRAY, bold=True, align=PP_ALIGN.CENTER)
            _add_picture(slide, path, 0.9, y, 12.15, 1.05, compress=True)


def _profile_fit_slides(presentation: Presentation, timepoint: int, paths: list[Path]) -> None:
    paths = sorted(paths, key=_z_number)
    for offset in range(0, len(paths), 2):
        chunk = paths[offset:offset + 2]
        slide = _blank_slide(presentation)
        _text(slide, f"{_time_label(timepoint)} · normalized profiles and tanh fits",
              0.35, 0.12, 12.65, 0.45, size=21, bold=True, align=PP_ALIGN.CENTER)
        for index, path in enumerate(chunk):
            x = 0.25 + index * 6.55
            _text(slide, f"Z={_z_number(path)}", x, 0.67, 6.25, 0.35, size=15, color=GRAY, bold=True, align=PP_ALIGN.CENTER)
            _add_picture(slide, path, x, 1.02, 6.25, 6.15)


def create_step8_presentation(run_dir: Path) -> Path:
    step8 = run_dir / "step_08_selected_timepoints_all_z"
    stitched_root = step8 / "stitched_images"
    timepoint_dirs = sorted(path for path in stitched_root.glob("t[0-9][0-9][0-9]") if path.is_dir())
    if not timepoint_dirs:
        raise ValueError("No Step 8 stitched-image timepoint folders were found.")
    presentation = _new_presentation()
    _title_slide(
        presentation,
        "Selected timepoints across all Z planes",
        "Step 08 · corrected stitched images, normalized tanh fits, and signed-slope heatmaps",
    )
    _section_slide(presentation, "Part 1 · Corrected stitched images", "Merged, GFP, and Cy5 · five Z planes per slide")
    for timepoint_dir in timepoint_dirs:
        timepoint = int(timepoint_dir.name[1:])
        _section_slide(presentation, _time_label(timepoint), "Step 08 stitched images")
        merge = list(timepoint_dir.glob("*_GFP-Cy5_corrected_feathered_mosaic.png"))
        gfp = list(timepoint_dir.glob("*_GFP_corrected_feathered_mosaic_preview.png"))
        cy5 = list(timepoint_dir.glob("*_Cy5_corrected_feathered_mosaic_preview.png"))
        # Keep presentations reproducible for older contact-sheet runs.
        if not merge:
            merge = list(timepoint_dir.glob("*_GFP-Cy5_merge.png"))
            gfp = list(timepoint_dir.glob("*_GFP_positions_left-to-right_preview.png"))
            cy5 = list(timepoint_dir.glob("*_Cy5_positions_left-to-right_preview.png"))
        groups = (
            ("GFP–Cy5 merge", merge, DARK),
            ("GFP", gfp, GREEN),
            ("Cy5", cy5, MAGENTA),
        )
        for label, paths, color in groups:
            if len(paths) != 26:
                raise ValueError(f"Expected 26 {label} images for t{timepoint:03d}; found {len(paths)}.")
            _stitched_group_slides(presentation, timepoint, label, paths, color)

    _section_slide(presentation, "Part 2 · Normalized profiles and tanh fits", "Two Z planes per slide, grouped by timepoint")
    fit_root = step8 / "profile_fits"
    for timepoint_dir in timepoint_dirs:
        timepoint = int(timepoint_dir.name[1:])
        paths = list(fit_root.glob(f"*_t{timepoint:03d}_z*_GFP-Cy5_tanh-fit.png"))
        if len(paths) != 26:
            raise ValueError(f"Expected 26 profile fits for t{timepoint:03d}; found {len(paths)}.")
        _section_slide(presentation, _time_label(timepoint), "Step 08 normalized profiles and tanh fits")
        _profile_fit_slides(presentation, timepoint, paths)

    _section_slide(presentation, "Part 3 · Signed-slope heatmaps", "Rows are the five selected timepoints; columns are Z planes 1–26")
    slide = _blank_slide(presentation)
    _text(slide, "Signed tanh slope across timepoints and Z", 0.4, 0.15, 12.5, 0.45, size=23, bold=True, align=PP_ALIGN.CENTER)
    heatmaps = (
        (step8 / "2026_07_30_gradient_GFP_signed_slope_T-by-Z_heatmap.png", "GFP", GREEN, 0.25),
        (step8 / "2026_07_30_gradient_Cy5_signed_slope_T-by-Z_heatmap.png", "Cy5", MAGENTA, 6.75),
    )
    for path, label, color, x in heatmaps:
        _text(slide, label, x, 0.72, 6.25, 0.35, size=17, color=color, bold=True, align=PP_ALIGN.CENTER)
        _add_picture(slide, path, x, 1.08, 6.25, 5.95)

    output = run_dir / "2026_07_30_step08_all_z_images_profiles_heatmaps.pptx"
    presentation.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("run_dir", type=Path)
    args = parser.parse_args()
    created = [
        create_single_z_image_presentation(args.run_dir),
        create_single_z_comparison(args.run_dir),
    ]
    step8_stitched = args.run_dir / "step_08_selected_timepoints_all_z" / "stitched_images"
    if step8_stitched.is_dir():
        created.append(create_step8_presentation(args.run_dir))
    for path in created:
        print(f"Created: {path.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
